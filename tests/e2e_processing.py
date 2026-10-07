"""端到端验证：上传 → 向量化 → 状态 → 分块 → 级联删除 → 跨用户隔离"""

import asyncio
import io
import sys
import uuid

import asyncpg
import httpx
from docx import Document

BASE = "http://127.0.0.1:8000"
sys.path.insert(0, "src")
from rag.config import settings  # noqa: E402

PASS, FAIL = 0, 0


def check(name: str, cond: bool, detail: str = ""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  [PASS] {name}")
    else:
        FAIL += 1
        print(f"  [FAIL] {name} {detail}")


def make_docx() -> bytes:
    doc = Document()
    doc.add_heading("第一章 引言", level=1)
    for i in range(4):
        doc.add_paragraph(
            f"这是引言部分的第 {i+1} 段正文内容，用于测试分块算法在较长文本下的表现。"
            "RAG 系统需要将文档切分为合适大小的语义块，以便后续检索。"
        )
    doc.add_heading("1.1 研究背景", level=2)
    for i in range(3):
        doc.add_paragraph(
            f"研究背景第 {i+1} 段。向量检索依赖高质量的 embedding，"
            "而 embedding 的质量与分块策略密切相关。"
        )
    doc.add_heading("第二章 方法", level=1)
    for i in range(3):
        doc.add_paragraph(
            f"方法部分第 {i+1} 段。递归切分优先保证句子和段落的完整性，"
            "避免在语义中间截断文本。"
        )
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_pdf() -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    y = 800
    c.setFont("Helvetica-Bold", 18)
    c.drawString(50, y, "Chapter 1 Introduction"); y -= 30
    c.setFont("Helvetica", 10)
    for i in range(8):
        c.drawString(50, y, f"This is body paragraph {i+1} about retrieval augmented generation.")
        y -= 16
    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, y, "1.1 Background"); y -= 24
    c.setFont("Helvetica", 10)
    for i in range(6):
        c.drawString(50, y, f"Background paragraph {i+1} discussing embedding quality and chunking.")
        y -= 16
    c.showPage()
    c.save()
    return buf.getvalue()


async def main():
    async with httpx.AsyncClient(base_url=BASE, timeout=120.0) as c:
        # ---- 用户 1 注册 ----
        u1 = f"e2e_u1_{uuid.uuid4().hex[:8]}"
        r = await c.post("/api/v1/auth/register", json={"username": u1, "password": "Test1234!"})
        check("用户1注册", r.status_code in (200, 201), f"{r.status_code} {r.text[:200]}")
        u1_id = r.json()["id"]

        r = await c.post("/api/v1/auth/login", json={"username": u1, "password": "Test1234!"})
        check("用户1登录", r.status_code == 200, f"{r.status_code} {r.text[:200]}")

        # ---- 上传 DOCX ----
        files = {"file": ("test_chunking.docx", make_docx(),
                          "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        r = await c.post("/api/v1/documents/upload", files=files)
        check("上传 DOCX", r.status_code == 201, f"{r.status_code} {r.text[:300]}")
        doc = r.json()
        doc_id = doc["id"]
        check("上传后状态=not_processed", doc.get("processing_status") == "not_processed", str(doc))

        # ---- 触发处理 ----
        r = await c.post(f"/api/v1/documents/{doc_id}/process", json={})
        check("触发 process (202)", r.status_code == 202, f"{r.status_code} {r.text[:300]}")

        # ---- 轮询状态 ----
        status = None
        for _ in range(60):
            await asyncio.sleep(1)
            r = await c.get(f"/api/v1/documents/{doc_id}/status")
            if r.status_code != 200:
                check("查询状态", False, f"{r.status_code} {r.text[:200]}")
                return
            status = r.json()
            if status["processing_status"] in ("embedded", "failed"):
                break

        check("最终状态=embedded", status and status["processing_status"] == "embedded",
              f"status={status}")
        check("chunk_count>0", status and status["chunk_count"] > 0,
              f"chunk_count={status.get('chunk_count') if status else None}")
        check("embedded_at 非空", bool(status and status.get("embedded_at")))

        # ---- 分块列表 ----
        r = await c.get(f"/api/v1/documents/{doc_id}/chunks")
        check("获取分块列表", r.status_code == 200, f"{r.status_code} {r.text[:200]}")
        chunks = r.json()["chunks"] if r.status_code == 200 else []
        check("分块数量与 chunk_count 一致",
              len(chunks) == status["chunk_count"], f"{len(chunks)} vs {status['chunk_count']}")
        check("分块索引有序", [ch["chunk_index"] for ch in chunks] == list(range(len(chunks))))
        check("分块含 section_title", any(ch.get("section_title") for ch in chunks),
              f"titles={[ch.get('section_title') for ch in chunks][:5]}")
        check("分块内容非空", all(ch["content_preview"].strip() for ch in chunks))

        # ---- 文档列表含状态字段 ----
        r = await c.get("/api/v1/documents/")
        listed = next((d for d in r.json()["documents"] if d["id"] == doc_id), None)
        check("列表含 processing_status", listed and listed["processing_status"] == "embedded")
        check("列表含 chunk_count", listed and listed["chunk_count"] > 0)

        # ---- 重复触发（已 embedded，允许重新处理）----
        r = await c.post(f"/api/v1/documents/{doc_id}/process", json={})
        check("重新触发 process (202)", r.status_code == 202, f"{r.status_code}")
        # 等它跑完
        for _ in range(60):
            await asyncio.sleep(1)
            r = await c.get(f"/api/v1/documents/{doc_id}/status")
            if r.json()["processing_status"] in ("embedded", "failed"):
                break
        check("重新处理后仍 embedded", r.json()["processing_status"] == "embedded")

        # ---- 上传 PDF 并验证预处理 ----
        r = await c.post("/api/v1/documents/upload", files={
            "file": ("test_chunking.pdf", make_pdf(), "application/pdf")})
        check("上传 PDF", r.status_code == 201, f"{r.status_code} {r.text[:200]}")
        pdf_id = r.json()["id"]
        r = await c.post(f"/api/v1/documents/{pdf_id}/process", json={})
        check("PDF 触发 process", r.status_code == 202, f"{r.status_code}")
        pdf_status = None
        for _ in range(60):
            await asyncio.sleep(1)
            r = await c.get(f"/api/v1/documents/{pdf_id}/status")
            pdf_status = r.json()
            if pdf_status["processing_status"] in ("embedded", "failed"):
                break
        check("PDF 状态=embedded", pdf_status["processing_status"] == "embedded",
              f"{pdf_status}")
        check("PDF chunk_count>0", pdf_status["chunk_count"] > 0,
              f"chunks={pdf_status.get('chunk_count')}")
        r = await c.get(f"/api/v1/documents/{pdf_id}/chunks")
        pdf_chunks = r.json()["chunks"] if r.status_code == 200 else []
        check("PDF 分块可读取", len(pdf_chunks) == pdf_status["chunk_count"])
        r = await c.delete(f"/api/v1/documents/{pdf_id}")
        check("PDF 文档删除", r.status_code == 204, f"{r.status_code}")

        # ---- 数据库校验向量（使用 RLS 上下文）----
        from rag.db import get_user_connection, get_pool, close_pool
        await get_pool()
        async with get_user_connection(u1_id) as conn:
            n_with_vec = await conn.fetchval(
                "SELECT COUNT(*) FROM document_chunks WHERE document_id=$1 AND embedding IS NOT NULL",
                uuid.UUID(doc_id))
            dim = await conn.fetchval(
                "SELECT vector_dims(embedding) FROM document_chunks WHERE document_id=$1 LIMIT 1",
                uuid.UUID(doc_id))
        check("DB 中向量非空", n_with_vec == status["chunk_count"], f"{n_with_vec}")
        check("向量维度=1024", dim == 1024, f"dim={dim}")

        # ---- 跨用户隔离 ----
        u2 = f"e2e_u2_{uuid.uuid4().hex[:8]}"
        # 注意：register 会自动登录并写 cookie，必须用独立的 c2 客户端，
        # 否则会覆盖 user1 客户端的 cookie
        async with httpx.AsyncClient(base_url=BASE, timeout=120.0) as c2:
            await c2.post("/api/v1/auth/register", json={"username": u2, "password": "Test1234!"})
            await c2.post("/api/v1/auth/login", json={"username": u2, "password": "Test1234!"})
            r = await c2.get(f"/api/v1/documents/{doc_id}/status")
            check("跨用户 status → 404", r.status_code == 404, f"{r.status_code}")
            r = await c2.get(f"/api/v1/documents/{doc_id}/chunks")
            check("跨用户 chunks → 404", r.status_code == 404, f"{r.status_code}")
            r = await c2.post(f"/api/v1/documents/{doc_id}/process", json={})
            check("跨用户 process → 404", r.status_code == 404, f"{r.status_code}")
            r = await c2.delete(f"/api/v1/documents/{doc_id}")
            check("跨用户 delete → 404", r.status_code == 404, f"{r.status_code}")

        # ---- 删除文档 → 级联删除 chunks ----
        r = await c.delete(f"/api/v1/documents/{doc_id}")
        check("删除文档 (204)", r.status_code == 204, f"{r.status_code} body={r.text[:200]}")
        async with get_user_connection(u1_id) as conn:
            remaining = await conn.fetchval(
                "SELECT COUNT(*) FROM document_chunks WHERE document_id=$1", uuid.UUID(doc_id))
        check("chunks 级联删除", remaining == 0, f"remaining={remaining}")
        await close_pool()

    print(f"\n===== 结果: {PASS} passed, {FAIL} failed =====")
    return 0 if FAIL == 0 else 1


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))