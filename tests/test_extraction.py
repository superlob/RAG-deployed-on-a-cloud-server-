"""文本提取单元测试"""

from io import BytesIO

import pytest

from rag.processing.extraction import ExtractedBlock, ProcessingError, extract_text


# ==================== DOCX ====================


def _make_docx(paragraphs: list[tuple[str, str | None]]) -> bytes:
    """生成测试用 DOCX

    Args:
        paragraphs: [(text, style), ...]  style 如 'Heading 1', 'Normal', None
    """
    from docx import Document

    doc = Document()
    for text, style in paragraphs:
        para = doc.add_paragraph(text)
        if style:
            para.style = doc.styles[style]
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_simple_paragraphs():
    """普通段落提取"""
    data = _make_docx(
        [
            ("第一段正文内容", "Normal"),
            ("第二段正文内容", "Normal"),
        ]
    )
    blocks = extract_text("docx", data)

    assert len(blocks) == 2
    assert blocks[0].text == "第一段正文内容"
    assert blocks[0].heading_level is None
    assert blocks[1].text == "第二段正文内容"


def test_docx_headings_detected():
    """标题样式识别"""
    data = _make_docx(
        [
            ("第一章 概述", "Heading 1"),
            ("这是第一章的内容", "Normal"),
            ("1.1 背景介绍", "Heading 2"),
            ("这是背景内容", "Normal"),
        ]
    )
    blocks = extract_text("docx", data)

    assert blocks[0].heading_level == 1
    assert blocks[0].text == "第一章 概述"
    assert blocks[1].heading_level is None
    assert blocks[2].heading_level == 2
    assert blocks[2].text == "1.1 背景介绍"


def test_docx_skip_empty_paragraphs():
    """跳过空白段落"""
    data = _make_docx(
        [
            ("实际内容", "Normal"),
            ("", "Normal"),
            ("更多内容", "Normal"),
        ]
    )
    blocks = extract_text("docx", data)
    assert len(blocks) == 2


def test_docx_with_table():
    """表格内容提取为行文本"""
    from docx import Document

    doc = Document()
    doc.add_paragraph("表格前文本", style=doc.styles["Normal"])

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "B1"
    table.cell(1, 0).text = "A2"
    table.cell(1, 1).text = "B2"

    doc.add_paragraph("表格后文本", style=doc.styles["Normal"])

    buf = BytesIO()
    doc.save(buf)
    blocks = extract_text("docx", buf.getvalue())

    # 期望: [表格前文本, A1 | B1, A2 | B2, 表格后文本]
    assert len(blocks) >= 3
    assert blocks[0].text == "表格前文本"
    assert "A1" in blocks[1].text and "B1" in blocks[1].text
    assert "A2" in blocks[2].text and "B2" in blocks[2].text


# ==================== PDF ====================


def _make_pdf(pages: list[tuple[list[tuple[str, int]], list[tuple[str, int]]]]) -> bytes:
    """生成测试用 PDF

    Args:
        pages: 每页 = (标题列表[(text,font_size)], 正文列表[(text,font_size)])
    """
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4

    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)

    for page_titles, page_body in pages:
        y = 800
        for text, size in page_titles:
            c.setFont("Helvetica-Bold", size)
            c.drawString(50, y, text)
            y -= size + 6

        for text, size in page_body:
            c.setFont("Helvetica", size)
            c.drawString(50, y, text)
            y -= size + 4

            if y < 60:
                c.showPage()
                y = 800

        c.showPage()

    c.save()
    return buf.getvalue()


def test_pdf_body_text():
    """PDF 正文文本提取"""
    data = _make_pdf(
        [
            (
                [],  # no titles
                [
                    ("This is the first line of body text.", 10),
                    ("This is the second line of body text.", 10),
                    ("And the third line.", 10),
                ],
            )
        ]
    )
    blocks = extract_text("pdf", data)

    assert len(blocks) >= 3
    texts = " ".join(b.text for b in blocks)
    assert "first line" in texts
    assert "second line" in texts
    assert "third line" in texts
    # 所有正文块没有标题级别
    for b in blocks:
        assert b.heading_level is None


def test_pdf_heading_detection():
    """PDF 标题检测（大字号文本识别为标题）"""
    data = _make_pdf(
        [
            (
                [("Chapter 1 Introduction", 18)],
                [
                    ("This is the introduction body text.", 10),
                    ("Introduction continues here.", 10),
                ],
            ),
            (
                [("1.1 Background", 14)],
                [
                    ("Background analysis body text.", 10),
                ],
            ),
        ]
    )
    blocks = extract_text("pdf", data)

    headings = [b for b in blocks if b.heading_level is not None]
    # 至少应检测到 title（大字号）
    # 算法依赖字号分布，标题字号(18,14)与正文(10)差异较大，理应被识别。
    assert len(headings) >= 1, f"至少应检测到 1 个标题，实际 headings={headings}"


def test_pdf_page_numbers_tracked():
    """页码追踪验证"""
    data = _make_pdf(
        [
            ([], [("Page one content here.", 10)]),
            ([], [("Page two content here.", 10)]),
        ]
    )
    blocks = extract_text("pdf", data)

    pages = set(b.page_number for b in blocks)
    assert 1 in pages
    assert 2 in pages


def test_pdf_images_excluded():
    """图片不产生文本输出（pypdf 只提取文本层内容）"""
    from io import BytesIO

    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    # 生成包含「文本」和「图形」的 PDF（使用英文，因为 reportlab 内置字体不支持中文）
    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)

    c.setFont("Helvetica", 10)
    c.drawString(50, 800, "Visible text here")
    # 画一个矩形（图像/图形操作，不应产生文本）
    c.rect(50, 700, 100, 50, fill=1)
    c.showPage()
    c.save()

    blocks = extract_text("pdf", buf.getvalue())
    texts = [b.text for b in blocks]
    assert any("Visible text" in t for t in texts)
    # 图形不应产生文本
    assert len(blocks) == 1, f"图片不应产生文本块，实际 {len(blocks)}"


# ==================== DOC ====================


def test_doc_format_raises():
    """.doc 格式应抛出 ProcessingError"""
    with pytest.raises(ProcessingError, match=".doc"):
        extract_text("doc", b"fake data")


# ==================== 边界情况 ====================


def test_empty_pdf():
    """空 PDF 返回空列表"""
    from reportlab.pdfgen import canvas

    buf = BytesIO()
    c = canvas.Canvas(buf)
    c.showPage()
    c.save()

    blocks = extract_text("pdf", buf.getvalue())
    assert blocks == []


def test_empty_docx():
    """空 DOCX 返回空列表"""
    buf = BytesIO()
    from docx import Document

    Document().save(buf)
    blocks = extract_text("docx", buf.getvalue())
    assert blocks == []


def test_unsupported_type():
    """不支持的文件类型抛出异常"""
    with pytest.raises(ProcessingError, match="不支持"):
        extract_text("txt", b"plain text")