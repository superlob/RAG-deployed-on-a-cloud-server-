"""Chatting Interface Live 集成测试

真实外部依赖：
- 真实 PostgreSQL（云端）
- 真实 LLM API（gpt-6-luna，经代理）
- 真实检索（Retrieval 模块：pgvector + pg_trgm + RRF，查询 embedding 走 DashScope）

默认跳过：设置环境变量 RUN_LIVE_TESTS=1 才运行。
"""

from __future__ import annotations

import asyncio
import io
import json
import uuid

import httpx
import pytest
from docx import Document

from rag.db import get_user_connection
from rag.main import app

pytestmark = pytest.mark.asyncio(loop_scope="session")


# ---------------------------------------------------------------------------
# 辅助
# ---------------------------------------------------------------------------


async def make_client() -> httpx.AsyncClient:
    """ASGI 传输 + 同事件循环（跳过 lifespan，池懒初始化）"""
    transport = httpx.ASGITransport(app=app)
    return httpx.AsyncClient(transport=transport, base_url="http://test", timeout=120.0)


async def register_and_login(client: httpx.AsyncClient) -> tuple[dict, dict]:
    """注册并登录一个新用户，返回 (user, cookies 已由 client 保存)"""
    username = f"live_chat_{uuid.uuid4().hex[:10]}"
    password = "test-pass-123"

    r = await client.post("/api/v1/auth/register", json={"username": username, "password": password})
    assert r.status_code == 201, f"注册失败: {r.status_code} {r.text}"
    user = r.json()

    r = await client.post("/api/v1/auth/login", json={"username": username, "password": password})
    assert r.status_code == 200, f"登录失败: {r.status_code} {r.text}"
    return user, r.json()


async def parse_sse(response: httpx.Response) -> list[dict]:
    """从流式响应中解析 SSE 事件列表 [{event, data}]（忽略心跳注释行）"""
    events: list[dict] = []
    event_name = None
    data_lines: list[str] = []
    async for line in response.aiter_lines():
        if line.startswith(":"):
            continue  # 心跳注释
        if line.startswith("event: "):
            event_name = line[7:]
        elif line.startswith("data: "):
            data_lines.append(line[6:])
        elif line == "":
            if event_name is not None and data_lines:
                try:
                    data = json.loads("\n".join(data_lines))
                except json.JSONDecodeError:
                    data = {"_raw": "\n".join(data_lines)}
                events.append({"event": event_name, "data": data})
            event_name = None
            data_lines = []
    # 兜底：流结束但最后一个事件没有空行
    if event_name is not None and data_lines:
        try:
            data = json.loads("\n".join(data_lines))
        except json.JSONDecodeError:
            data = {"_raw": "\n".join(data_lines)}
        events.append({"event": event_name, "data": data})
    return events


async def get_message_rows(user_id: str, conversation_id: str) -> list:
    async with get_user_connection(user_id) as conn:
        return await conn.fetch(
            "SELECT role, content, tool_calls, sources, status, error "
            "FROM messages WHERE conversation_id = $1 ORDER BY seq",
            conversation_id,
        )


# ---------------------------------------------------------------------------
# 测试
# ---------------------------------------------------------------------------


async def test_live_crud_via_http():
    """HTTP CRUD：注册/登录/创建/列表/重命名/删除会话（不走 LLM）"""
    async with await make_client() as client:
        user, _ = await register_and_login(client)
        uid = user["id"]

        # 创建
        r = await client.post("/api/v1/conversations", json={})
        assert r.status_code == 201, r.text
        conv = r.json()
        assert conv["title"] == "新对话"
        cid = conv["id"]

        # 列表
        r = await client.get("/api/v1/conversations")
        assert r.status_code == 200
        listing = r.json()
        assert listing["total"] >= 1
        assert any(c["id"] == cid for c in listing["conversations"])

        # 重命名
        r = await client.patch(f"/api/v1/conversations/{cid}", json={"title": "重命名后的会话"})
        assert r.status_code == 200
        assert r.json()["title"] == "重命名后的会话"

        # 空历史
        r = await client.get(f"/api/v1/conversations/{cid}/messages")
        assert r.status_code == 200
        assert r.json()["total"] == 0

        # 不存在的会话 → 404
        r = await client.get(f"/api/v1/conversations/{uuid.uuid4()}/messages")
        assert r.status_code == 404

        # 删除
        r = await client.delete(f"/api/v1/conversations/{cid}")
        assert r.status_code == 204
        r = await client.get(f"/api/v1/conversations/{cid}/messages")
        assert r.status_code == 404


async def test_live_sse_plain_greeting():
    """真实 LLM：问候语直接回答（无工具调用），SSE 流 + 持久化均正确"""
    async with await make_client() as client:
        user, _ = await register_and_login(client)
        uid = user["id"]

        r = await client.post("/api/v1/conversations", json={"title": "live-plain"})
        cid = r.json()["id"]

        events: list[dict] = []
        async with client.stream(
            "POST", f"/api/v1/conversations/{cid}/messages",
            json={"content": "你好，请用一句话简短介绍你自己。"},
        ) as resp:
            assert resp.status_code == 200
            assert resp.headers["content-type"].startswith("text/event-stream")
            events = await parse_sse(resp)

        types = [e["event"] for e in events]
        assert "status" in types
        assert "delta" in types, f"缺少 delta 事件: {types}"
        assert "done" == types[-1], f"流必须以 done 结束: {types}"
        assert "error" not in types

        # 拼接出的回复非空
        text = "".join(e["data"]["content"] for e in events if e["event"] == "delta")
        assert len(text.strip()) > 0

        # 问候语不应触发检索
        assert "tool_call" not in types

        # 持久化：user + assistant
        rows = await get_message_rows(uid, cid)
        assert [r["role"] for r in rows] == ["user", "assistant"]
        assert rows[1]["status"] == "completed"
        assert rows[1]["content"].strip() == text.strip()

        # 自动标题（首条消息）
        r = await client.get("/api/v1/conversations")
        titles = {c["id"]: c["title"] for c in r.json()["conversations"]}
        assert titles[cid].startswith("你好，请用一句话")


def make_conclusion_docx() -> bytes:
    """生成一篇含「核心结论」的测试 DOCX，供真实检索流验证"""
    doc = Document()
    doc.add_heading("第一章 引言", level=1)
    doc.add_paragraph(
        "本文研究检索增强生成（RAG）系统。RAG 通过将文档切分为语义块并向量化存储，"
        "在回答问题时先检索相关片段再生成答案。"
    )
    doc.add_heading("第二章 核心结论", level=1)
    doc.add_paragraph(
        "核心结论：将文档切分为高质量的语义块能够显著提升检索准确率和回答质量。"
    )
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def test_live_sse_retrieval_flow():
    """真实 LLM + 真实检索：上传并向量化文档 → 提问 → 检索 → 带来源回答"""
    async with await make_client() as client:
        user, _ = await register_and_login(client)
        uid = user["id"]

        # 上传并向量化一篇含「核心结论」的文档
        files = {
            "file": (
                "retrieval_live.docx",
                make_conclusion_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        r = await client.post("/api/v1/documents/upload", files=files)
        assert r.status_code == 201, r.text
        doc_id = r.json()["id"]

        r = await client.post(f"/api/v1/documents/{doc_id}/process", json={})
        assert r.status_code == 202, r.text

        status = None
        for _ in range(90):
            await asyncio.sleep(1)
            r = await client.get(f"/api/v1/documents/{doc_id}/status")
            assert r.status_code == 200, r.text
            status = r.json()
            if status["processing_status"] in ("embedded", "failed"):
                break
        assert status and status["processing_status"] == "embedded", f"处理失败: {status}"

        r = await client.post("/api/v1/conversations", json={"title": "live-retrieval"})
        assert r.status_code == 201, r.text
        cid = r.json()["id"]

        events: list[dict] = []
        async with client.stream(
            "POST", f"/api/v1/conversations/{cid}/messages",
            json={"content": "我上传的论文里，核心结论是什么？"},
        ) as resp:
            assert resp.status_code == 200
            events = await parse_sse(resp)

        types = [e["event"] for e in events]
        assert "done" == types[-1]
        assert "error" not in types, f"出现错误事件: {events}"

        # 文档类问题应触发真实检索并返回来源
        assert "tool_call" in types, f"未触发检索工具: {types}"
        assert "sources" in types
        src = next(e for e in events if e["event"] == "sources")
        sources = src["data"]["sources"]
        assert len(sources) > 0
        assert "retrieval_live.docx" in {s.get("document_title") for s in sources}

        # 最终回答非空
        text = "".join(e["data"]["content"] for e in events if e["event"] == "delta")
        assert len(text.strip()) > 0

        # 持久化：user → assistant(tool_calls) → tool → assistant(sources)
        rows = await get_message_rows(uid, cid)
        roles = [r["role"] for r in rows]
        assert roles == ["user", "assistant", "tool", "assistant"]
        final = rows[3]
        assert final["status"] == "completed"
        stored_sources = final["sources"]
        if isinstance(stored_sources, str):
            stored_sources = json.loads(stored_sources)
        assert len(stored_sources) > 0


async def test_live_multi_turn_history():
    """真实 LLM：两轮对话，第二轮基于上下文（引用第一轮内容）"""
    async with await make_client() as client:
        user, _ = await register_and_login(client)
        uid = user["id"]

        r = await client.post("/api/v1/conversations", json={"title": "live-multi"})
        cid = r.json()["id"]

        # 第一轮
        async with client.stream(
            "POST", f"/api/v1/conversations/{cid}/messages",
            json={"content": "请记住：我的幸运数字是 42。"},
        ) as resp:
            assert resp.status_code == 200
            await parse_sse(resp)

        # 第二轮（引用第一轮）
        events: list[dict] = []
        async with client.stream(
            "POST", f"/api/v1/conversations/{cid}/messages",
            json={"content": "我的幸运数字是多少？"},
        ) as resp:
            assert resp.status_code == 200
            events = await parse_sse(resp)

        assert "done" == [e["event"] for e in events][-1]
        text = "".join(e["data"]["content"] for e in events if e["event"] == "delta")
        assert "42" in text, f"第二轮未引用第一轮内容: {text[:200]}"

        # 持久化：4 条消息（2 轮 × user+assistant）
        rows = await get_message_rows(uid, cid)
        assert len(rows) == 4
        assert rows[0]["content"] == "请记住：我的幸运数字是 42。"
